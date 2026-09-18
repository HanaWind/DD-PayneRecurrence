from pathlib import Path
import json
import math

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DOCS = ROOT / "docs"
DOCS.mkdir(exist_ok=True)
ASSETS = DOCS / "report_assets"
ASSETS.mkdir(exist_ok=True)
OUT_DOCX = DOCS / "DD-Payne_复现工作说明.docx"
OUT_PROMPT = DOCS / "DD-Payne_组会PPT提示词.txt"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", sz="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), sz)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Aptos", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, space_after=5, line_spacing=1.15):
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing


def add_text(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    style_paragraph(p)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    style_paragraph(p, space_after=2)
    r = p.add_run(text)
    set_run_font(r, size=10.2)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, size=15 if level == 1 else 11.5, bold=True)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        if widths:
            cell.width = Inches(widths[i])
        set_cell_shading(cell, "274C77")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, space_after=0, line_spacing=1.0)
        r = p.add_run(str(text))
        set_run_font(r, size=font_size, bold=True, color="FFFFFF")
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            if widths:
                cell.width = Inches(widths[i])
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2 == 1:
                set_cell_shading(cell, "F2F5F8")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, space_after=0, line_spacing=1.05)
            r = p.add_run(str(text))
            set_run_font(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_after=6, line_spacing=1.0)
    r = p.add_run(text)
    set_run_font(r, size=9, color="555555")
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color="666666")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color="666666")


def make_smoke_plot():
    metrics = []
    with (ROOT / "outputs/smoke/run/metrics.jsonl").open(encoding="utf-8") as f:
        for line in f:
            metrics.append(json.loads(line))
    steps = [m["step"] for m in metrics]
    train = [m["train_reconstruction"] for m in metrics]
    val = [m.get("validation_reconstruction") for m in metrics]
    grad_steps = [m["step"] for m in metrics if m.get("train_gradient") is not None]
    grad = [m["train_gradient"] for m in metrics if m.get("train_gradient") is not None]
    width, height = 1200, 520
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    left, top, right, bottom = 90, 70, 1130, 440
    d.rectangle((left, top, right, bottom), outline="#9CA3AF", width=2)
    d.text((left, 22), "Smoke test learning curve", fill="#111827")
    vmax = max(train + [v for v in val if v is not None]) * 1.05
    vmin = 0
    def xy(s, v):
        x = left + (s - min(steps)) / (max(steps) - min(steps)) * (right-left)
        y = bottom - (v-vmin)/(vmax-vmin) * (bottom-top)
        return (int(x), int(y))
    for ytick in range(0, 4):
        value = vmax * ytick / 3
        y = bottom - ytick / 3 * (bottom-top)
        d.line((left, y, right, y), fill="#E5E7EB", width=1)
        d.text((12, int(y)-8), f"{value:.0f}", fill="#4B5563")
    d.line([xy(s, v) for s, v in zip(steps, train)], fill="#274C77", width=4)
    val_x = [s for s, v in zip(steps, val) if v is not None]
    val_y = [v for v in val if v is not None]
    d.line([xy(s, v) for s, v in zip(val_x, val_y)], fill="#D97706", width=4)
    for s, v in zip(steps, train):
        x, y = xy(s, v); d.ellipse((x-4,y-4,x+4,y+4), fill="#274C77")
    for s, v in zip(val_x, val_y):
        x, y = xy(s, v); d.rectangle((x-4,y-4,x+4,y+4), fill="#D97706")
    d.text((left, bottom+18), "step", fill="#4B5563")
    d.text((right-250, 24), "blue: train   orange: validation", fill="#4B5563")
    path = ASSETS / "smoke_learning_curve.png"
    img.save(path)
    return path


def build_prompt():
    prompt = """你是一名擅长天文学术汇报设计的PPT生成AI。请根据下面的项目材料，生成一份用于组会汇报的中文PPT，建议 12-14 页，风格简洁、理工科、以深蓝和灰白为主色，避免商业宣传感。听众是熟悉光谱分析但不一定了解 DD-Payne 细节的研究组成员。

汇报主题：基于 LAMOST DR9 低分辨率光谱的 DD-Payne 复现进展

汇报目标：
1. 说明我已经完成了哪些代码和工程复现工作。
2. 用一条清晰链路解释“观测光谱 -> 预处理 -> 标签到光谱前向模型 -> 训练 -> 反演标签 -> 质量控制”。
3. 讲清楚当前结果能证明什么、还不能证明什么。
4. 补充我对微湍流、自转与宏观湍流展宽、视向速度、分辨率与线扩散函数的理解。
5. 梳理噪声来源谱系，并在每类噪声旁边标注处理方式。
6. 最后一页给出下一阶段可执行路线，而不是笼统地说“继续优化”。

必须准确使用的项目事实：
- 当前仓库面向 LAMOST DR9 v2.0 LRS，文档记录本地 spectra 子集有 125,075 条 fits.gz 光谱，来自 14 个观测夜，单个 COADD 抽样文件约 3909 像素。
- 当前本地数据没有 APOGEE/GALAH 高分辨率监督标签，也没有 Kurucz/ATLAS12/SYNTHE 梯度谱，因此目前完成的是可运行的复现基础设施和合成 smoke test，不应表述为已经得到论文级 15 维元素丰度结果。
- 代码已实现：FITS 读取、质量掩码、静止系校正、50 A 高斯伪连续谱归一化、公共对数波长网格重采样、恒星级分组切分、逐像素 Payne 网络、观测谱加权重建损失、有限差分物理梯度正则、检查点、标签拟合、Fisher 不确定度接口。
- 15 维标签为 teff、logg、vmic、fe_h、c_fe、n_fe、o_fe、mg_fe、al_fe、si_fe、ca_fe、ti_fe、cr_fe、mn_fe、ni_fe。
- 当前预处理配置：3800-8800 A，log10 波长步长 0.0001；g 波段 S/N 下限 20；ANDMASK 非零像素拒绝；有效像素比例至少 0.80；按 source_id 做 80/10/10 训练/验证/测试划分，seed=42。
- Payne 网络为每个输出像素独立参数的两隐层 MLP，每层 40 个隐藏单元；正式配置的 batch size=128，最多 10000 steps，学习率从 0.01 指数退火到 0.0001，pixel_chunk_size=512。
- smoke test 使用合成数据和梯度库，CPU 配置 20 steps，训练重建损失从约 3196 降到约 923，验证重建损失从约 2957 降到约 1099；这只能证明训练路径、梯度正则、验证和检查点能够工作。

页面结构要求：
第1页 标题页：项目名、数据域、汇报目标；配一条从光谱到标签的简化流程箭头。
第2页 先给结论：已完成的工程模块、当前最大缺口、当前结果的证据等级。用“已完成 / 可运行验证 / 尚未具备科学训练条件”三栏。
第3页 为什么是 DD-Payne：说明传统数据驱动模型与物理梯度约束的区别，突出“标签到光谱”的前向模型和反向标签拟合。
第4页 端到端流程：FITS -> mask -> 静止系 -> 50 A 伪连续谱 -> log-lambda 重采样 -> 训练/验证/测试 -> inference。
第5页 数据与预处理细节：波长范围、S/N、质量掩码、source_id 分组切分，解释每个设置会影响什么。
第6页 模型结构：画出 15 维标签输入、两个 40 单元隐藏层、每像素独立输出的示意；标注 pixel chunk 节省显存。
第7页 损失函数：重建损失、有限差分梯度正则、标签权重；给出公式的直观解释，不要堆太多数学。
第8页 当前 smoke test 证据：折线图展示训练/验证重建损失下降，旁边写清“合成数据、不是科学结果”。
第9页 参数理解一：微湍流 vmic；解释未分辨小尺度速度场、对饱和线和等效宽度的影响，并给一个高 vmic/低 vmic 的例子。
第10页 参数理解二：自转与宏观湍流展宽；对比旋转核和宏观湍流展宽，说明两者都能让线变宽但物理形状和参数退化不同。
第11页 参数理解三：视向速度、分辨率与 LSF；给出 Δλ/λ≈v/c、R=λ/Δλ，并用 5000 A、R≈1800 的例子说明约 2.8 A 的分辨率尺度；强调训练前后 LSF 必须一致。
第12页 噪声来源谱系：按“观测随机噪声 / 数据处理残差 / 仪器与标定系统误差 / 模型与标签系统误差”分组，用表格列出来源、光谱表现、处理方式。
第13页 当前缺口与风险：监督标签、物理梯度库、LSF 随波长/光纤变化、RV 误差、标签系统误差、训练域外推；每项对应下一步动作。
第14页 下一阶段路线和总结：先拿 APOGEE DR17 做主训练标签，GALAH DR3 做外部验证；生成与 LAMOST 网格和 LSF 一致的梯度库；先小样本过拟合再正式训练；用重复观测、星团、外部巡天、梯度一致性和 Cramer-Rao 做验证。

视觉与表达要求：
- 每页最多 4-6 个要点，优先图、流程图和小表格，避免大段文字。
- 术语首次出现时给中文解释，括号保留英文缩写，例如线扩散函数（LSF）、视向速度（RV）、分辨率 R。
- 所有“结果”都标注证据等级：代码单元测试、合成 smoke test、真实科学训练、外部验证。
- 不要虚构论文级精度、元素丰度散点或外部验证结果。
- 结尾用一句话总结：当前工作已经打通复现链路，但科学结论还必须建立在监督标签和物理梯度库补齐之后。

请输出：
1. 每页的标题。
2. 每页 3-6 条可直接放到幻灯片上的中文要点。
3. 每页推荐的图或示意图。
4. 一份 8-10 分钟中文讲稿，按页分段，语气像组会汇报而不是论文摘要。
5. 对“当前不能声称什么”单独给出一页脚注式提醒。"""
    OUT_PROMPT.write_text(prompt, encoding="utf-8")


def build_docx():
    smoke_plot = make_smoke_plot()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for sname in ("Title", "Heading 1", "Heading 2"):
        styles[sname].font.name = "Aptos Display"
        styles[sname]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles[sname].font.color.rgb = RGBColor(0, 0, 0)

    footer = sec.footer.paragraphs[0]
    add_page_number(footer)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run("DD Payne 复现工作说明")
    set_run_font(r, name="Aptos Display", size=23, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(subtitle, space_after=16, line_spacing=1.0)
    r = subtitle.add_run("面向 LAMOST DR9 低分辨率光谱的工程复现与方法理解")
    set_run_font(r, size=12, color="4B5563")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(meta, space_after=18, line_spacing=1.0)
    r = meta.add_run("组会汇报配套工作文档  2026 年 9 月")
    set_run_font(r, size=9.5, color="6B7280")

    add_heading(doc, "摘要", 1)
    add_text(doc, "本项目依据 DD-Payne 和 The Payne 的公开方法描述，独立重建了面向 LAMOST DR9 低分辨率光谱的可运行基础框架。当前已经打通 FITS 数据读取、质量掩码、静止系校正、伪连续谱归一化、公共波长网格重采样、恒星级数据切分、逐像素 Payne 前向模型、观测谱重建损失、物理梯度正则、检查点保存、标签反演和 Fisher 不确定度估计接口。")
    add_text(doc, "需要明确的是，当前本地光谱目录没有 APOGEE/GALAH 高分辨率监督标签，也没有与 LAMOST 波长网格、连续谱定义和线扩散函数一致的 Kurucz/ATLAS12/SYNTHE 梯度库。因此现阶段的证据是“工程链路已实现并通过测试与合成 smoke test”，还不是“已经得到论文级科学丰度结果”。")

    add_heading(doc, "结论先行", 1)
    add_table(doc, ["状态", "已经完成", "当前含义"], [
        ("工程实现", "LAMOST 读取、预处理、模型、损失、训练、推断接口和配置文件", "代码路径完整，可继续接入真实监督数据"),
        ("可运行验证", "单元测试与合成 smoke test；20 steps 中训练损失约 3196 降至 923", "证明数值路径、梯度正则、验证和检查点能够工作"),
        ("尚未具备", "APOGEE/GALAH 标签、物理梯度库、真实训练集和外部验证", "不能声称已有可信的 15 维元素丰度测量"),
    ], widths=[1.0, 3.0, 3.0])

    add_heading(doc, "一 项目目标与复现边界", 1)
    add_text(doc, "DD-Payne 的核心是学习从恒星标签到归一化光谱的前向映射，再在观测谱上反向拟合标签。与只依赖标签共变关系的数据驱动模型相比，物理梯度谱正则把每个标签引起的光谱变化方向加入训练约束，降低模型把元素间天体物理相关性误当成谱线信息的风险。")
    add_text(doc, "本仓库针对 LAMOST DR9 LRS 建立了独立实现，而不是声称复刻作者内部代码。仓库文档记录本地 spectra 子集约有 125,075 条 fits.gz 光谱，来自 14 个观测夜；抽样 COADD 扩展约有 3909 个像素。当前数据只适合完成数据工程、接口联调和方法验证，正式科学训练还需要补齐高分辨率标签和物理梯度库。")

    add_heading(doc, "二 已完成的实现链路", 1)
    add_table(doc, ["阶段", "实现内容", "关键设置或输出"], [
        ("数据读取", "读取 LAMOST fits.gz 的 COADD 扩展，获得 flux、ivar、wavelength、ANDMASK、ORMASK 和 normalization", "保留 obsid、source_id、坐标、S/N、红移等元数据"),
        ("质量控制", "星类和 S/N 筛选；ANDMASK 非零像素默认拒绝；无效像素的 ivar 置零", "最低 g 波段 S/N=20，有效像素比例至少 0.80"),
        ("预处理", "按 RV 变换到静止系；50 A 高斯伪连续谱归一化；排除强 Balmer/DIB 区域；重采样到公共 log-lambda 网格", "3800-8800 A，log10 步长 0.0001"),
        ("数据切分", "按 source_id 分组切分，避免同一颗星的重复观测泄漏", "train/validation/test=80/10/10，seed=42"),
        ("前向模型", "每个像素独立参数的两隐层 Payne MLP", "15 标签；hidden size=40；pixel chunk=512"),
        ("训练与反演", "加权重建损失 + 有限差分梯度正则；Adam 训练；逐星标签拟合和 Fisher 不确定度", "正式配置最多 10000 steps，学习率 0.01 -> 0.0001"),
    ], widths=[1.0, 3.5, 2.5])

    add_heading(doc, "三 模型与损失的理解", 1)
    add_text(doc, "模型输入是 15 维标签，输出是每个波长像素上的归一化光谱值。网络对每个像素保留独立的权重张量，因此能表达不同波长处不同的标签响应；训练时按像素块计算，主要目的是降低大波长网格下的显存峰值。")
    add_text(doc, "观测谱重建项可理解为按像素逆方差加权的平方残差：高质量像素权重大，坏像素或无覆盖像素不参与损失。梯度正则项用有限差分估计模型对每个标签的光谱导数，再与物理梯度库比较；当前配置给 Teff、logg、vmic 较小权重，给 Fe/H 和元素丰度较大权重，具体数值仍应通过交叉验证调整。")
    add_table(doc, ["模块", "直观问题", "当前实现"], [
        ("重建损失", "模型生成的光谱是否贴近观测谱", "weighted_reconstruction_sum，使用 ivar 加权"),
        ("物理梯度正则", "标签变化引起的谱形变化方向是否物理合理", "finite_difference_gradient_sum，与梯度库逐标签比较"),
        ("反演", "给定一条观测谱，什么标签能生成最接近的光谱", "Adam 优化标签，输出 reduced chi2、有效像素、边界标志"),
        ("不确定度", "当前光谱的信息量足以把标签定到什么程度", "Fisher 矩阵伪逆，输出标签标准差估计"),
    ], widths=[1.3, 3.2, 2.5])

    add_heading(doc, "四 当前验证证据", 1)
    add_text(doc, "当前 smoke test 使用合成光谱和合成梯度库，配置为 CPU、20 steps、16 条 batch、6 个隐藏单元。它不是 LAMOST 科学结果，而是用于验证训练循环、梯度正则、验证损失、检查点保存和加载路径的最小闭环。")
    doc.add_picture(str(smoke_plot), width=Inches(6.35))
    add_caption(doc, "图 1  合成 smoke test 的训练与验证重建损失；下降趋势证明训练路径可运行，不代表真实标签精度。")
    add_table(doc, ["证据", "观测到的结果", "证据等级"], [
        ("单元测试", "覆盖波长网格、连续谱掩码、坏像素权重、模型梯度、梯度损失和 Fisher 不确定度", "代码级"),
        ("CPU smoke test", "训练重建损失约 3196 -> 923；验证重建损失约 2957 -> 1099；梯度正则项可计算", "合成数据级"),
        ("真实科学训练", "尚未执行；缺监督标签和物理梯度库", "未完成"),
        ("外部验证", "尚未执行；需重复观测、星团和外部高分辨率巡天", "未完成"),
    ], widths=[1.4, 4.2, 1.4])

    add_heading(doc, "五 关键光谱参数的理解", 1)
    add_table(doc, ["参数", "物理含义", "会影响什么", "例子"], [
        ("微湍流 vmic", "未被模型网格显式解析的小尺度随机速度场，通常以 Doppler 展宽进入谱线形成", "尤其影响饱和线的线芯、等效宽度和由线强反推的丰度", "在相同 Fe/H 下，vmic 增大可让饱和线变宽、减弱饱和效应；若忽略它，可能把线强差异误判成丰度差异"),
        ("自转 v sin i", "恒星表面不同投影速度的积分结果，形成旋转核卷积", "线轮廓展宽和形状；高速自转会混合相邻吸收线", "v sin i 增大时，窄线变宽并出现更明显的旋转型轮廓，弱线可能被淹没"),
        ("宏观湍流", "较大尺度的对流或表面速度场，常用径向-切向等核近似", "也会展宽谱线，但轮廓形状与旋转不同；与 v sin i 存在退化", "低分辨率下两者都可能表现为线宽增加，需要高 S/N 或外部先验区分"),
        ("视向速度 RV", "沿视线方向的整体速度，使谱线发生多普勒平移", "线的位置、静止系对齐和重采样误差", "5000 A 处 RV=30 km/s 约对应 0.5 A 位移；不校正会让模型把移位当成参数变化"),
        ("分辨率 R", "R=lambda/Delta lambda，描述仪器分开相近谱线的能力", "线宽、线混合和可辨识的元素信息量", "R≈1800 且 lambda=5000 A 时，Delta lambda≈2.8 A；很多细线会被混合"),
        ("线扩散函数 LSF", "仪器对窄线响应的完整核，通常随波长、光纤和观测条件变化", "模型合成谱与观测谱的线宽和线形是否一致", "梯度谱若未用同一 LSF 卷积，物理梯度方向会被错误的线宽差污染"),
    ], widths=[1.1, 2.0, 2.0, 1.9], font_size=8.7)
    add_text(doc, "这几类参数在低分辨率光谱中常存在退化：分辨率和 LSF 先决定仪器能看到多细的线形，RV 决定线是否对齐，vmic、旋转和宏观湍流则共同改变线宽或线芯。因而训练集、物理梯度库和观测谱必须使用同一波长定义、LSF 和归一化规则。")

    add_heading(doc, "六 噪声来源谱系与处理方式", 1)
    add_table(doc, ["噪声类别", "来源", "光谱表现", "处理方式"], [
        ("观测随机噪声", "光子计数噪声、读出噪声、暗电流、增益不确定度", "像素间随机起伏，低通量处相对噪声变大", "用 ivar 进入加权损失；质量筛选设 S/N 下限；必要时做仪器标定和噪声传播"),
        ("天空与背景残差", "天空发射线减除不完全、散射光、背景估计偏差", "局部尖峰、宽肩或基线起伏，常集中在特定波段", "质量位屏蔽；在连续谱估计中排除问题区间；对残差波段做分箱质检"),
        ("坏像素与宇宙线", "探测器坏列、饱和、宇宙线、抽取失败", "孤立异常点或整段无效像素", "ANDMASK 非零默认拒绝；无效像素 ivar=0；重采样时保持覆盖掩码"),
        ("波长与 RV 误差", "波长标定漂移、管线 RV 偏差、静止系变换误差", "整条谱或局部线系发生小幅错位", "按 RV 变换到静止系；抽查高 S/N 线位；必要时联合拟合 RV 或加 RV 质量位"),
        ("连续谱与通量标定", "响应曲线误差、宽尺度通量偏差、伪连续谱估计偏差", "吸收线深度和宽尺度斜率被系统改变", "50 A 高斯伪连续谱；强 Balmer/DIB 区间排除；保留 pipeline normalization 作为可选模式"),
        ("LSF 与仪器系统误差", "LSF 随波长、光纤、板和观测条件变化，平均 LSF 近似不足", "线宽和线形系统性偏差，可能被误认为 vmic 或旋转", "梯度谱按同一 LSF 卷积；按观测域做 LSF 质检；正式模型考虑分域或 LSF 条件化"),
        ("标签与模型系统误差", "APOGEE/GALAH 标签零点、1D LTE、大气模型、线表和训练域外推", "跨巡天偏差、参数边界堆积、某些元素残差成片", "统一标签标尺；物理梯度正则；外部巡天和星团验证；输出边界/外推标志"),
    ], widths=[1.25, 2.05, 2.0, 1.7], font_size=8.45)
    add_text(doc, "处理原则可以概括为：随机噪声进入权重，明确坏数据进入掩码，波长问题先对齐，连续谱问题先统一定义，LSF 和标签系统误差则必须通过物理建模与外部验证识别，不能仅靠提高网络容量消除。")

    add_heading(doc, "七 当前缺口与风险", 1)
    add_table(doc, ["缺口或风险", "为什么重要", "下一步动作"], [
        ("缺高分辨率监督标签", "没有元素丰度真值锚点，无法训练或评估 15 维标签", "优先接入 APOGEE DR17；保留来源、误差、质量位和版本"),
        ("缺物理梯度库", "数据驱动模型可能利用元素共变关系而非真实谱线响应", "用 ATLAS12/SYNTHE 或验证过的等价物，统一 LSF、网格、静止系和连续谱"),
        ("LSF 变化", "线宽误差会与 vmic、旋转和宏观湍流退化", "建立按波长/光纤/观测域的 LSF 诊断，必要时条件化建模"),
        ("训练域不足", "当前本地子集覆盖可能不均，域外标签会被错误当作正常测量", "绘制标签凸包和 S/N 覆盖；边界样本标记或拒绝"),
        ("RV 与重复观测", "错位或泄漏会夸大模型表现", "抽查线位；按 source_id 分组；用重复观测测一致性"),
    ], widths=[1.65, 3.1, 2.25])

    add_heading(doc, "八 下一阶段执行路线", 1)
    for item in [
        "第一阶段：完成 APOGEE DR17 主训练标签的交叉匹配和质量筛选，统一 15 维标签命名、太阳标尺和误差字段。",
        "第二阶段：制作与 LAMOST 训练网格一致的物理梯度库，逐标签有限差分，并通过同一 LSF、静止系、连续谱和掩码处理。",
        "第三阶段：用 32-128 颗训练星和少量梯度参考点做过拟合检查，确认重建下降、梯度方向一致、检查点可复现。",
        "第四阶段：扩大到正式训练集，扫描梯度权重、连续谱宽度和参考点批量；保存配置、环境、数据版本和检查点哈希。",
        "第五阶段：进行重复观测一致性、开放星团/球状星团散布、GALAH 外部比较、逐元素梯度一致性和 Fisher/Cramer-Rao 分析。",
        "第六阶段：固定检查点后再做全量 LAMOST 推断，输出标签、误差、reduced chi2、有效像素数、边界/外推标志和训练覆盖距离。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "九 组会汇报时需要主动说明的边界", 1)
    add_text(doc, "可以声称：复现代码框架已经形成，关键数据处理和模型接口已经实现，单元测试和合成 smoke test 已经通过，工程上可以进入真实监督数据接入阶段。")
    add_text(doc, "不能声称：当前已经获得可靠的 LAMOST 15 维元素丰度目录，已经达到论文精度，已经完成 APOGEE/GALAH 外部验证，或已经证明某一组超参数在科学上最优。")
    add_text(doc, "一句话总结：当前工作已经打通 DD-Payne 的复现链路，但科学结论还必须建立在监督标签、物理梯度库、LSF 一致性和独立验证全部补齐之后。")

    add_heading(doc, "附录 项目中可直接引用的配置与文件", 1)
    add_table(doc, ["内容", "位置"], [
        ("项目说明和复现边界", "README.md；docs/REPRODUCTION_GUIDE_ZH.md"),
        ("完整后续执行清单", "STEPS.md"),
        ("预处理配置", "configs/preprocess_lamost.yaml"),
        ("训练配置", "configs/train_lamost.yaml"),
        ("推断配置", "configs/infer_lamost.yaml"),
        ("核心预处理", "src/ddpayne/data/preprocess.py"),
        ("Payne 网络", "src/ddpayne/models/payne.py"),
        ("损失和梯度库", "src/ddpayne/models/losses.py"),
        ("标签拟合和 Fisher 不确定度", "src/ddpayne/inference/fit.py"),
        ("合成 smoke test 结果", "outputs/smoke/run/metrics.jsonl、best.pt、last.pt"),
    ], widths=[2.5, 4.5])

    doc.core_properties.title = "DD Payne 复现工作说明"
    doc.core_properties.subject = "LAMOST DR9 低分辨率光谱 DD-Payne 复现"
    doc.core_properties.author = ""
    doc.core_properties.comments = ""
    doc.save(OUT_DOCX)


if __name__ == "__main__":
    build_prompt()
    build_docx()
    print(OUT_DOCX)
    print(OUT_PROMPT)
